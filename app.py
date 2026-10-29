import streamlit as st
from groq import Groq
from docx import Document
from PIL import Image
import io
import base64

st.set_page_config(page_title="Digitalizador de Informes", layout="centered")
st.title("Digitalizador Automatico de Informes")

if "GROQ_API_KEY" in st.secrets:
    api_key = st.secrets["GROQ_API_KEY"].strip()
else:
    api_key = st.sidebar.text_input("Ingresa tu Groq API Key:", type="password").strip()

if not api_key:
    st.info("Por favor, configura la API Key de Groq para continuar.")
    st.stop()

client = Groq(api_key=api_key)

archivos = st.file_uploader("Selecciona la(s) imagen(es)", type=["jpg", "png", "jpeg"], accept_multiple_files=True)

if archivos and st.button("Procesar y Generar Word", type="primary"):
    doc = Document()
    doc.add_heading("Informe Digitalizado", level=1)
    contenido_extraido = False

    for idx, archivo in enumerate(archivos, start=1):
        try:
            # Convertir imagen a Base64 para enviarla a Groq
            bytes_data = archivo.getvalue()
            base64_image = base64.b64encode(bytes_data).decode('utf-8')

            response = client.chat.completions.create(
                model="llama-3.2-11b-vision-preview",
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": "Transcribe y organiza toda la informacion de esta imagen en formato Markdown estructurado (# para titulos, - para listas)."},
                            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}}
                        ]
                    }
                ]
            )

            texto_resultado = response.choices[0].message.content

            if texto_resultado:
                contenido_extraido = True
                lineas = texto_resultado.split('\n')
                for linea in lineas:
                    l = linea.strip()
                    if l.startswith('# '):
                        doc.add_heading(l.replace('# ', ''), level=1)
                    elif l.startswith('## '):
                        doc.add_heading(l.replace('## ', ''), level=2)
                    elif l.startswith('- ') or l.startswith('* '):
                        doc.add_paragraph(l[2:], style='List Bullet')
                    elif l:
                        doc.add_paragraph(l)

        except Exception as e:
            st.error(f"Error procesando {archivo.name}: {e}")

    if contenido_extraido:
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        st.success("Documento procesado con exito.")
        st.download_button("Descargar Word (.docx)", data=buffer, file_name="Informe.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")